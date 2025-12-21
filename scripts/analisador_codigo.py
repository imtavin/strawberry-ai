import ast
import os
from docx import Document
from docx.shared import Pt

def analisar_arquivo(caminho_arquivo):
    """Analisa a estrutura de um arquivo Python e retorna um dicionário com suas informações."""
    with open(caminho_arquivo, "r", encoding="utf-8") as f:
        codigo = f.read()

    arvore = ast.parse(codigo)
    resultado = {
        "arquivo": os.path.basename(caminho_arquivo),
        "imports": [],
        "funcoes": [],
        "classes": [],
        "variaveis": []
    }

    for no in arvore.body:
        if isinstance(no, ast.Import):
            for alias in no.names:
                resultado["imports"].append(alias.name)
        elif isinstance(no, ast.ImportFrom):
            mod = no.module if no.module else ""
            for alias in no.names:
                resultado["imports"].append(f"{mod}.{alias.name}")
        elif isinstance(no, ast.FunctionDef):
            doc = ast.get_docstring(no)
            parametros = [arg.arg for arg in no.args.args]
            resultado["funcoes"].append({
                "nome": no.name,
                "linha": no.lineno,
                "parametros": parametros,
                "docstring": doc
            })
        elif isinstance(no, ast.ClassDef):
            metodos = []
            for elem in no.body:
                if isinstance(elem, ast.FunctionDef):
                    doc_m = ast.get_docstring(elem)
                    parametros_m = [arg.arg for arg in elem.args.args]
                    metodos.append({
                        "nome": elem.name,
                        "linha": elem.lineno,
                        "parametros": parametros_m,
                        "docstring": doc_m
                    })
            resultado["classes"].append({
                "nome": no.name,
                "linha": no.lineno,
                "docstring": ast.get_docstring(no),
                "metodos": metodos
            })
        elif isinstance(no, ast.Assign):
            for alvo in no.targets:
                if isinstance(alvo, ast.Name):
                    resultado["variaveis"].append(alvo.id)

    return resultado


def gerar_relatorio_docx(dados, nome_arquivo="relatorio_codigo.docx"):
    """Gera um relatório DOCX a partir dos dados estruturais do código."""
    doc = Document()
    doc.add_heading('Relatório de Análise Estrutural do Código', 0)

    for arquivo in dados:
        doc.add_heading(f"Arquivo: {arquivo['arquivo']}", level=1)

        # Importações
        if arquivo["imports"]:
            doc.add_heading("Importações:", level=2)
            for imp in arquivo["imports"]:
                doc.add_paragraph(f"- {imp}")

        # Variáveis globais
        if arquivo["variaveis"]:
            doc.add_heading("Variáveis Globais:", level=2)
            for var in arquivo["variaveis"]:
                doc.add_paragraph(f"- {var}")

        # Funções
        if arquivo["funcoes"]:
            doc.add_heading("Funções:", level=2)
            for func in arquivo["funcoes"]:
                p = doc.add_paragraph(f"{func['nome']} (linha {func['linha']})\n")
                p.add_run(f"Parâmetros: {', '.join(func['parametros']) or 'nenhum'}\n").italic = True
                if func["docstring"]:
                    p.add_run(f"Descrição: {func['docstring']}\n")

        # Classes
        if arquivo["classes"]:
            doc.add_heading("Classes:", level=2)
            for cls in arquivo["classes"]:
                p = doc.add_paragraph(f"{cls['nome']} (linha {cls['linha']})\n")
                if cls["docstring"]:
                    p.add_run(f"Descrição: {cls['docstring']}\n")
                if cls["metodos"]:
                    doc.add_paragraph("Métodos:")
                    for met in cls["metodos"]:
                        sub = doc.add_paragraph(f"- {met['nome']} (linha {met['linha']})", style='List Bullet')
                        sub.add_run(f"\nParâmetros: {', '.join(met['parametros']) or 'nenhum'}").italic = True
                        if met["docstring"]:
                            sub.add_run(f"\nDescrição: {met['docstring']}")

        doc.add_page_break()

    doc.save(nome_arquivo)
    print(f"Relatório gerado com sucesso: {nome_arquivo}")


def analisar_diretorio(pasta):
    """Percorre todos os arquivos .py da pasta e gera relatório."""
    dados = []
    for raiz, _, arquivos in os.walk(pasta):
        for nome in arquivos:
            if nome.endswith(".py"):
                caminho = os.path.join(raiz, nome)
                print(f"Analisando: {caminho}")
                dados.append(analisar_arquivo(caminho))
    gerar_relatorio_docx(dados)


if __name__ == "__main__":
    pasta = input("Digite o caminho da pasta do projeto: ").strip()
    analisar_diretorio(pasta)
